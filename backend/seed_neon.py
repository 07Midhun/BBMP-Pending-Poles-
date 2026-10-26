import os
import sys
import datetime
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(BASE_DIR))

from dotenv import load_dotenv
if (BASE_DIR / ".env").exists():
    load_dotenv(dotenv_path=BASE_DIR / ".env")
load_dotenv()

from database import engine, SessionLocal, init_db
from models_db import PoleDB
from excel_parser import generate_sample_schnell_iot_data, classify_old_lamp, normalize_lamp_type
from sqlalchemy.dialects.postgresql import insert

def load_installed_pole_numbers():
    installed_file = BASE_DIR / "Parameterised_Lamp_Installation_Report_Schnelliot_Live_Lamp_Installation_Report_Pivot table (1).xlsx"
    if not installed_file.is_file():
        print("No installation report Excel found, skipping installed set.", flush=True)
        return set()
    
    print(f"Reading installation report from {installed_file.name}...", flush=True)
    import openpyxl
    wb = openpyxl.load_workbook(str(installed_file), read_only=True)
    sheet = wb.active
    rows = sheet.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    
    pole_idx = header.index("Pole ID") if "Pole ID" in header else 1
    installed = set()
    for r in rows:
        val = r[pole_idx] if pole_idx < len(r) else None
        if val:
            installed.add(str(val).strip().upper())
            
    print(f"Loaded {len(installed)} installed pole IDs from installation report.", flush=True)
    return installed

def seed_database():
    if engine is None or SessionLocal is None:
        print("ERROR: DATABASE_URL is not set or invalid in environment.", flush=True)
        return False

    print("Initializing Neon database schema...", flush=True)
    init_db()

    session = SessionLocal()
    try:
        existing_count = session.query(PoleDB).count()
        print(f"Current poles count in Neon: {existing_count}", flush=True)

        installed_set = load_installed_pole_numbers()
        excel_path = BASE_DIR / "Pole_survey_report(BBMP).xlsx"

        if excel_path.is_file():
            print(f"Streaming master survey poles from {excel_path.name} with low memory...", flush=True)
            import openpyxl
            wb = openpyxl.load_workbook(str(excel_path), read_only=True)
            sheet = wb.active
            rows = sheet.iter_rows(values_only=True)
            header = [str(h).strip() if h is not None else "" for h in next(rows)]

            col_region = header.index("Region") if "Region" in header else 0
            col_zone = header.index("Zone") if "Zone" in header else 1
            col_ward = header.index("Ward") if "Ward" in header else 2
            col_pole = header.index("Pole ID") if "Pole ID" in header else 3
            col_lat = header.index("Latitude") if "Latitude" in header else 14
            col_lng = header.index("Longitude") if "Longitude" in header else 15
            col_lamp = header.index("Lamp Type") if "Lamp Type" in header else 25

            batch = {}
            batch_size = 1000
            total_inserted = 0

            for row in rows:
                if not row or len(row) <= col_pole:
                    continue
                pole_no = str(row[col_pole] or "").strip().upper()
                if not pole_no or len(pole_no) > 64:
                    continue

                raw_lamp = str(row[col_lamp] or "").strip() if col_lamp < len(row) else ""
                norm_lamp = normalize_lamp_type(raw_lamp)
                old_lamp = classify_old_lamp(norm_lamp)

                try:
                    lat = float(row[col_lat] or 0.0) if col_lat < len(row) else 0.0
                except (ValueError, TypeError):
                    lat = 0.0

                try:
                    lng = float(row[col_lng] or 0.0) if col_lng < len(row) else 0.0
                except (ValueError, TypeError):
                    lng = 0.0

                reg = str(row[col_region] or "EAST").strip().upper()[:64] if col_region < len(row) else "EAST"
                zn = str(row[col_zone] or "").strip()[:64] if col_zone < len(row) else ""
                wd = str(row[col_ward] or "").strip()[:64] if col_ward < len(row) else ""

                batch[pole_no] = {
                    "pole_number": pole_no,
                    "region": reg,
                    "zone": zn,
                    "ward": wd,
                    "pole_old_lamp": old_lamp[:32],
                    "lamp_type": norm_lamp[:64],
                    "latitude": lat,
                    "longitude": lng,
                    "is_installed": pole_no in installed_set,
                    "uploaded_slots": [],
                }

                if len(batch) >= batch_size:
                    records = list(batch.values())
                    stmt = insert(PoleDB).values(records)
                    stmt = stmt.on_conflict_do_update(
                        index_elements=["pole_number"],
                        set_={
                            "region": stmt.excluded.region,
                            "zone": stmt.excluded.zone,
                            "ward": stmt.excluded.ward,
                            "pole_old_lamp": stmt.excluded.pole_old_lamp,
                            "lamp_type": stmt.excluded.lamp_type,
                            "latitude": stmt.excluded.latitude,
                            "longitude": stmt.excluded.longitude,
                            "is_installed": stmt.excluded.is_installed,
                        }
                    )
                    session.execute(stmt)
                    session.commit()
                    total_inserted += len(records)
                    now_str = datetime.datetime.now().strftime("%H:%M:%S")
                    print(f"  [{now_str}] Seeded {total_inserted} poles into Neon...", flush=True)
                    batch.clear()

            if batch:
                records = list(batch.values())
                stmt = insert(PoleDB).values(records)
                stmt = stmt.on_conflict_do_update(
                    index_elements=["pole_number"],
                    set_={
                        "region": stmt.excluded.region,
                        "zone": stmt.excluded.zone,
                        "ward": stmt.excluded.ward,
                        "pole_old_lamp": stmt.excluded.pole_old_lamp,
                        "lamp_type": stmt.excluded.lamp_type,
                        "latitude": stmt.excluded.latitude,
                        "longitude": stmt.excluded.longitude,
                        "is_installed": stmt.excluded.is_installed,
                    }
                )
                session.execute(stmt)
                session.commit()
                total_inserted += len(records)
                print(f"  Final batch committed. Total seeded: {total_inserted} poles.", flush=True)

        else:
            print("No Excel survey file found. Generating default BBMP dataset...", flush=True)
            iot_data = generate_sample_schnell_iot_data()
            poles_to_insert = iot_data["master_poles"]
            installed_set = set(iot_data["installed_report"])
            records = []
            for p in poles_to_insert:
                pole_no = str(p.get("pole_number", "")).strip().upper()
                if not pole_no:
                    continue
                records.append({
                    "pole_number": pole_no,
                    "region": str(p.get("region", "EAST")).strip().upper(),
                    "zone": str(p.get("zone", "")).strip(),
                    "ward": str(p.get("ward", "")).strip(),
                    "pole_old_lamp": str(p.get("pole_old_lamp", "Empty")).strip(),
                    "lamp_type": str(p.get("lamp_type", "-")).strip(),
                    "latitude": float(p.get("latitude", 0.0)),
                    "longitude": float(p.get("longitude", 0.0)),
                    "is_installed": pole_no in installed_set,
                    "uploaded_slots": [],
                })
            stmt = insert(PoleDB).values(records)
            session.execute(stmt)
            session.commit()

        final_count = session.query(PoleDB).count()
        pending_count = session.query(PoleDB).filter(PoleDB.is_installed == False).count()
        installed_count = session.query(PoleDB).filter(PoleDB.is_installed == True).count()
        print(f"\n==================================================", flush=True)
        print(f"SUCCESS! Neon PostgreSQL Seeding Completed:", flush=True)
        print(f"  Total Master Poles   : {final_count}", flush=True)
        print(f"  Installed Poles      : {installed_count}", flush=True)
        print(f"  Pending Poles for App: {pending_count}", flush=True)
        print(f"==================================================\n", flush=True)
        return True
    except Exception as exc:
        session.rollback()
        print(f"ERROR while seeding database: {exc}", flush=True)
        return False
    finally:
        session.close()

if __name__ == "__main__":
    seed_database()
